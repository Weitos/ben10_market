import logging
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import create_engine, Column, Integer, String, Float, ForeignKey
from sqlalchemy.orm import declarative_base, sessionmaker, relationship

logger = logging.getLogger(__name__)

Base = declarative_base()

# Оригинальные HEX-цвета для openpyxl
GREEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
RED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")


class ProductModel(Base):
    __tablename__ = 'products'

    id = Column(Integer, primary_key=True)
    name = Column(String, nullable=False)
    barcode = Column(String, unique=True, nullable=False)

    metrics = relationship("FinancialMetricModel", uselist=False, back_populates="product", cascade="all, delete-orphan")


class FinancialMetricModel(Base):
    __tablename__ = 'financial_metrics'

    product_id = Column(Integer, ForeignKey('products.id'), primary_key=True)
    price = Column(Float, nullable=False)
    cost_price = Column(Float, nullable=False)
    commission_pct = Column(Float, nullable=False)
    net_profit = Column(Float)
    margin_pct = Column(Float)

    product = relationship("ProductModel", back_populates="metrics")


class DatabaseClient:
    def __init__(self, db_url: str = None):
        """Инициализация базы данных SQLite."""
        self.db_url = db_url or os.getenv('DATABASE_URL', 'sqlite:///data/market_data.db')
        
        # Создаем директорию для БД, если её нет
        db_path = self.db_url.replace('sqlite:///', '')
        if '/' in db_path:
            os.makedirs(os.path.dirname(db_path), exist_ok=True)
            
        self.engine = create_engine(self.db_url)
        self.Session = sessionmaker(bind=self.engine)
        Base.metadata.create_all(self.engine)
        logger.info(f"База данных инициализирована: {self.db_url}")

    def save_market_data(self, df: pd.DataFrame):
        """
        Сохраняет рассчитанный в analytics.py DataFrame в две связанные таблицы.
        """
        session = self.Session()
        try:
            for _, row in df.iterrows():
                # 1. Добавляем или обновляем сам продукт
                product = session.query(ProductModel).filter_by(id=int(row['id'])).first()
                if not product:
                    product = ProductModel(
                        id=int(row['id']),
                        name=str(row['name']),
                        barcode=str(row['barcode'])
                    )
                    session.add(product)
                else:
                    product.name = str(row['name'])
                    product.barcode = str(row['barcode'])

                # 2. Добавляем или обновляем финансовые метрики (с результатами из analytics.py)
                metrics = session.query(FinancialMetricModel).filter_by(product_id=product.id).first()
                if not metrics:
                    metrics = FinancialMetricModel(product_id=product.id)
                    session.add(metrics)

                metrics.price = float(row['price'])
                metrics.cost_price = float(row['cost_price'])
                metrics.commission_pct = float(row['commission_pct'])
                metrics.net_profit = round(float(row['net_margin_rub']), 2)  # Колонка из вашего analytics.py
                metrics.margin_pct = round(float(row['margin_pct']), 2)      # Колонка из вашего analytics.py

            session.commit()
            logger.info(f"Успешно сохранено/обновлено продуктов в БД: {len(df)}")
        except Exception as e:
            session.rollback()
            logger.error(f"Ошибка сохранения в БД: {e}")
            raise
        finally:
            session.close()

    def export_to_excel(self, file_path: str = "data/Calculated_Market_Data.xlsx"):
        """
        Загружает данные из БД через JOIN, сортирует и сохраняет в Excel с подсветкой ТОП-3.
        """
        try:
            query = """
                SELECT p.id, p.name, p.barcode, f.price, f.cost_price, f.commission_pct, f.net_profit, f.margin_pct
                FROM products p
                JOIN financial_metrics f ON p.id = f.product_id
            """
            df = pd.read_sql_query(query, self.engine)

            if df.empty:
                logger.warning("Нет данных в БД для экспорта.")
                return

            # Сортируем по профиту сверху вниз, как в analytics.py
            df = df.sort_values(by="net_profit", ascending=False).reset_index(drop=True)
            df_values = df.values.tolist()

            wb = Workbook()
            ws = wb.active
            ws.title = "Рейтинг товаров"

            # Заголовки
            headers = ['ID', 'Название', 'Штрихкод', 'Цена', 'Себестоимость', 'Комиссия %', 'Чистая прибыль (руб)', 'Маржинальность %']
            ws.append(headers)

            for row in df_values:
                ws.append(row)

            total_rows = len(df_values)
            max_col = len(headers)

            # Расчет индексов строк для openpyxl (+2 из-за шапки и 1-индексации)
            top_best_rows = [i + 2 for i in range(min(3, total_rows))]
            top_worst_rows = [i + 2 for i in range(max(0, total_rows - 3), total_rows)]

            # Применяем стили к строкам
            for r_idx in top_best_rows:
                for c_idx in range(1, max_col + 1):
                    ws.cell(row=r_idx, column=c_idx).fill = GREEN_FILL

            for r_idx in top_worst_rows:
                if r_idx not in top_best_rows:
                    for c_idx in range(1, max_col + 1):
                        ws.cell(row=r_idx, column=c_idx).fill = RED_FILL

            # Автоматическое выравнивание ширины колонок (чтобы текст не обрезался)
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            wb.save(file_path)
            logger.info(f"Отчет Excel успешно сгенерирован: {file_path}")

        except Exception as e:
            logger.error(f"Ошибка экспорта в Excel: {e}")
            raise
